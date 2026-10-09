"""Importação em lote: CSV, XLSX, texto colado e cadastro manual.

Fluxo: parse -> sugestão de mapeamento -> prévia com validação -> confirmação.
Nenhuma linha é descartada em silêncio: duplicadas e inválidas entram no
relatório da importação (e linhas inválidas são gravadas com seus erros).
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Correspondence, DeliveryBatch
from .normalize import clean_spaces, norm_token_text, normalize_address, parse_address_line

FIELDS = [
    "external_id", "recipient", "street", "number", "complement", "neighborhood",
    "city", "cep", "notes", "full_address",
]
SYNONYMS = {
    "external_id": ["id", "identificador", "codigo", "cod", "objeto", "registro", "rastreio", "nº objeto", "n objeto"],
    "recipient": ["destinatario", "nome", "cliente", "recebedor"],
    "street": ["logradouro", "rua", "endereco", "end", "via"],
    "number": ["numero", "num", "n", "nº", "no"],
    "complement": ["complemento", "compl", "apto", "apartamento"],
    "neighborhood": ["bairro"],
    "city": ["municipio", "cidade", "localidade"],
    "cep": ["cep", "codigo postal"],
    "notes": ["observacoes", "observacao", "obs", "nota", "notas"],
    "full_address": ["endereco completo", "endereço completo", "endereco_completo"],
}

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000


class ImportError_(ValueError):
    pass


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[list[str]]
    original_text: str
    sha256: str
    detected_delimiter: str | None = None
    sheets: list[str] = field(default_factory=list)
    has_header: bool = True


def _decode(content: bytes, encoding: str | None) -> str:
    if encoding:
        return content.decode(encoding)
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportError_("Não foi possível decodificar o arquivo.")


def parse_delimited(text: str, delimiter: str | None = None, has_header: bool = True) -> ParsedTable:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln for ln in text.split("\n") if ln.strip()]
    if not lines:
        raise ImportError_("Conteúdo vazio.")
    if delimiter in (None, "", "auto"):
        # ";", tab e "|" raramente aparecem em endereços; "," só é aceita como
        # separador com cabeçalho e contagem constante (senão: uma coluna por linha).
        sample = lines[:30]
        delimiter = None
        for d in (";", "\t", "|", ","):
            counts = {ln.count(d) for ln in sample}
            if min(counts) > 0 and (d != "," or (has_header and len(counts) == 1)):
                delimiter = d
                break
    if delimiter is None:
        rows = [[ln] for ln in lines]  # uma coluna: endereço completo por linha
    else:
        rows = [[clean_spaces(c) for c in r] for r in csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)]
    if has_header:
        headers, rows = rows[0], rows[1:]
    else:
        width = max(len(r) for r in rows)
        headers = [f"coluna_{i + 1}" for i in range(width)]
    if len(rows) > MAX_ROWS:
        raise ImportError_(f"Limite de {MAX_ROWS} linhas por lote.")
    headers = [clean_spaces(h) or f"coluna_{i + 1}" for i, h in enumerate(headers)]
    return ParsedTable(headers, rows, text, hashlib.sha256(text.encode()).hexdigest(), delimiter,
                       has_header=has_header)


def parse_xlsx(content: bytes, sheet: str | None = None, has_header: bool = True) -> ParsedTable:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_(f"XLSX inválido: {exc}") from exc
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(values_only=True):
        vals = ["" if v is None else clean_spaces(str(v)) for v in r]
        if any(vals):
            rows.append(vals)
        if len(rows) > MAX_ROWS + 1:
            raise ImportError_(f"Limite de {MAX_ROWS} linhas por lote.")
    if not rows:
        raise ImportError_("Planilha vazia.")
    if has_header:
        headers, rows = rows[0], rows[1:]
    else:
        headers = [f"coluna_{i + 1}" for i in range(max(len(r) for r in rows))]
    headers = [h or f"coluna_{i + 1}" for i, h in enumerate(headers)]
    # Entrada original preservada como CSV (;) das células lidas.
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(headers)
    w.writerows(rows)
    return ParsedTable(headers, rows, buf.getvalue(), hashlib.sha256(content).hexdigest(),
                       None, [s.title for s in wb.worksheets], has_header)


def parse_upload(content: bytes, filename: str, options: dict) -> ParsedTable:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportError_("Arquivo maior que 10 MB.")
    name = (filename or "").lower()
    has_header = options.get("has_header", True)
    if name.endswith(".xlsx") or content[:2] == b"PK":
        if not (name.endswith(".xlsx") or name == ""):
            raise ImportError_("Tipo de arquivo não suportado.")
        return parse_xlsx(content, options.get("sheet"), has_header)
    if name and not name.endswith((".csv", ".txt", ".tsv")):
        raise ImportError_("Envie CSV, TXT ou XLSX.")
    text = _decode(content, options.get("encoding"))
    if "\x00" in text:
        raise ImportError_("Arquivo binário não suportado.")
    return parse_delimited(text, options.get("delimiter"), has_header)


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {f: None for f in FIELDS}
    normed = {h: norm_token_text(h).lower() for h in headers}
    for f in FIELDS:
        for h, n in normed.items():
            if h in mapping.values():
                continue
            if n == f or n in SYNONYMS[f] or n.replace(" ", "_") in SYNONYMS[f]:
                mapping[f] = h
                break
    if len(headers) == 1 and not mapping["street"]:
        mapping["full_address"] = headers[0]
    return mapping


@dataclass
class MappedRow:
    row_number: int
    raw: dict
    values: dict
    errors: list[str]
    warnings: list[str]
    address_key: str | None


def map_rows(table: ParsedTable, mapping: dict, default_city: str = "Santos") -> list[MappedRow]:
    idx = {h: i for i, h in enumerate(table.headers)}
    unknown = [v for v in mapping.values() if v and v not in idx]
    if unknown:
        raise ImportError_(f"Colunas inexistentes no mapeamento: {unknown}")
    if not mapping.get("street") and not mapping.get("full_address"):
        raise ImportError_("Mapeie a coluna de logradouro ou de endereço completo.")
    out = []
    for n, row in enumerate(table.rows, start=2 if table.has_header else 1):
        raw = {h: (row[i] if i < len(row) else "") for h, i in idx.items()}
        vals = {f: clean_spaces(raw[mapping[f]]) if mapping.get(f) else None for f in FIELDS}
        if vals.get("full_address"):
            parsed = parse_address_line(vals["full_address"])
            for k, v in parsed.items():
                if not vals.get(k):
                    vals[k] = v
        vals = {k: (v or None) for k, v in vals.items()}
        if not any(vals.values()):
            continue  # linha totalmente vazia não é uma correspondência
        na = normalize_address(vals["street"], vals["number"], vals["neighborhood"],
                               vals["city"], vals["cep"], default_city=default_city)
        out.append(MappedRow(n, raw, vals, list(na.errors), list(na.issues),
                             na.key if not na.errors else None))
    return out


def preview(table: ParsedTable, mapping: dict | None, limit: int = 30) -> dict:
    mapping = mapping or suggest_mapping(table.headers)
    rows = map_rows(table, mapping)
    ids = [r.values.get("external_id") for r in rows if r.values.get("external_id")]
    dup_in_file = sorted({i for i in ids if ids.count(i) > 1})
    return {
        "headers": table.headers,
        "mapping": mapping,
        "delimiter": table.detected_delimiter,
        "sheets": table.sheets,
        "total_rows": len(rows),
        "error_rows": sum(1 for r in rows if r.errors),
        "warning_rows": sum(1 for r in rows if r.warnings),
        "duplicate_ids_in_file": dup_in_file,
        "unique_addresses": len({r.address_key for r in rows if r.address_key}),
        "rows": [
            {"row_number": r.row_number, "values": {k: v for k, v in r.values.items() if k != "recipient"},
             "has_recipient": bool(r.values.get("recipient")), "errors": r.errors, "warnings": r.warnings}
            for r in rows[:limit]
        ],
    }


def commit_import(db: Session, org_id: uuid.UUID, user_id: uuid.UUID | None, table: ParsedTable,
                  mapping: dict, source_type: str, reference: str, service_date: date,
                  filename: str | None, options: dict, allow_same_file: bool = False) -> DeliveryBatch:
    existing = db.scalar(select(DeliveryBatch).where(
        DeliveryBatch.org_id == org_id, DeliveryBatch.content_sha256 == table.sha256,
        DeliveryBatch.status != "failed"))
    if existing and not allow_same_file:
        raise ImportError_(
            f"Este conteúdo já foi importado no lote '{existing.reference}'. "
            "Confirme explicitamente para importar novamente."
        )
    rows = map_rows(table, mapping)
    batch = DeliveryBatch(
        org_id=org_id, reference=reference, service_date=service_date, source_type=source_type,
        original_filename=filename, content_sha256=table.sha256, original_content=table.original_text,
        options=options, column_mapping=mapping, status="importing", created_by=user_id,
    )
    db.add(batch)
    db.flush()

    ext_ids = [r.values["external_id"] for r in rows if r.values.get("external_id")]
    already = set()
    if ext_ids:
        already = set(db.scalars(select(Correspondence.external_id).where(
            Correspondence.org_id == org_id, Correspondence.external_id.in_(ext_ids))))
    seen: set[str] = set()
    report = {"inserted": 0, "with_errors": [], "duplicates": [], "warnings": 0}
    objs = []
    for r in rows:
        ext = r.values.get("external_id")
        if ext and (ext in already or ext in seen):
            # mesma carta (mesmo identificador): não duplica, mas registra.
            report["duplicates"].append({"row_number": r.row_number, "external_id": ext,
                                         "reason": "identificador já importado"})
            continue
        if ext:
            seen.add(ext)
        objs.append(Correspondence(
            org_id=org_id, batch_id=batch.id, row_number=r.row_number, external_id=ext,
            raw_row=r.raw, recipient=r.values.get("recipient"), street=r.values.get("street"),
            number=r.values.get("number"), complement=r.values.get("complement"),
            neighborhood=r.values.get("neighborhood"), city=r.values.get("city"),
            cep=r.values.get("cep"), notes=r.values.get("notes"),
            validation_errors=[*({"level": "error", "msg": e} for e in r.errors),
                               *({"level": "warning", "msg": w} for w in r.warnings)],
            status="pending",
        ))
        if r.errors:
            report["with_errors"].append({"row_number": r.row_number, "errors": r.errors})
        if r.warnings:
            report["warnings"] += 1
    db.add_all(objs)
    report["inserted"] = len(objs)
    batch.row_count = len(rows)
    batch.error_count = len(report["with_errors"])
    batch.import_report = report
    batch.status = "imported"
    db.flush()
    return batch


def manual_table(entries: list[dict]) -> ParsedTable:
    headers = [f for f in FIELDS if f != "full_address"]
    rows = [[clean_spaces(str(e.get(h) or "")) for h in headers] for e in entries]
    text = "\n".join(";".join(r) for r in [headers, *rows])
    return ParsedTable(headers, rows, text,
                       hashlib.sha256((text + str(uuid.uuid4())).encode()).hexdigest(), None,
                       has_header=False)
